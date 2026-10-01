# fuel_tracking/services/cph_abaque_import.py
"""
Import de l'abaque CPH GE PRP 50 Hz (ABAQUE_CPH_GE_PRP_50HZ.xlsx), partagé par
la commande import_cph_abaque et l'upload depuis l'onglet Contrôle CPH :
  - feuille « Abaque CPH »          → CphCurve (upsert par ID courbe)
  - feuille « Mappage inventaire »  → CphInventoryMapping (upsert par libellé + kVA inventaire)

Les validations métier déjà saisies sont conservées ; une validation de mappage
dont la courbe n'est plus candidate est annulée et signalée. Le cos φ déjà
ajusté sur une courbe existante n'est pas écrasé. a, b, c sont recalculés à
partir des 3 points 50/75/100 % et tout écart > 0,05 est signalé (l'abaque fait
foi, rien n'est corrigé en silence).
"""
from decimal import Decimal

import openpyxl
from django.db import transaction
from django.utils import timezone

from fuel_tracking.models import CphCurve, CphInventoryMapping
from fuel_tracking.services.cph_engine import normalize_label
from fuel_tracking.services.cph_matching import apply_auto_matching

CURVE_COLUMNS = {
    "curve_id": "ID courbe", "manufacturer": "Fabricant", "model": "Modèle", "model_key": "Clé modèle",
    "variant": "Liste/variante", "prp_kva": "PRP kVA", "prp_kw": "PRP kW", "power_factor": "cos φ",
    "voltage_v": "Tension (V)", "phases": "Phases", "frequency": "Fréquence", "regime": "Régime",
    "conso_25_l_h": "Conso 25 % extrapolée (L/h)", "conso_50_l_h": "Conso 50 % (L/h)",
    "conso_75_l_h": "Conso 75 % (L/h)", "conso_100_l_h": "Conso 100 % (L/h)",
    "coef_a": "a", "coef_b": "b", "coef_c": "c", "domain": "Domaine", "status": "Statut",
    "source": "Source", "source_url": "URL source", "note": "Note",
}
MAPPING_COLUMNS = {
    "inventory_label": "Libellé inventaire", "inventory_kva": "kVA inventaire", "site_count": "Nombre de sites",
    "normalized_key": "Clé normalisée", "candidate_curve_ids": "IDs courbe candidats",
    "candidate_models": "Modèles candidats", "abaque_status": "Statut", "action_required": "Action requise",
}
DECIMAL_FIELDS = {"prp_kva", "prp_kw", "power_factor", "conso_25_l_h", "conso_50_l_h", "conso_75_l_h",
                  "conso_100_l_h", "coef_a", "coef_b", "coef_c", "inventory_kva"}
INT_FIELDS = {"voltage_v", "phases", "site_count"}


def _read_sheet(wb, name, columns):
    if name not in wb.sheetnames:
        raise ValueError(f"Feuille « {name} » absente du fichier.")
    rows = list(wb[name].iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [c for c in columns.values() if c not in header]
    if missing:
        raise ValueError(f"Feuille « {name} » : colonnes manquantes {missing}")
    idx = {field: header.index(col) for field, col in columns.items()}
    out = []
    for r in rows[1:]:
        if not r or r[idx[next(iter(columns))]] in (None, ""):
            continue
        rec = {}
        for field, i in idx.items():
            v = r[i] if i < len(r) else None
            if isinstance(v, str):
                v = v.strip()
            if field in DECIMAL_FIELDS:
                v = Decimal(str(v)) if v not in (None, "") else None
            elif field in INT_FIELDS:
                v = int(v) if v not in (None, "") else None
            elif v is None:
                v = ""
            rec[field] = v
        out.append(rec)
    return out


def fitted_coefficients(y50: Decimal, y75: Decimal, y100: Decimal) -> tuple[Decimal, Decimal, Decimal]:
    a = (y100 - 2 * y75 + y50) * 8
    b = (y100 - y50) * 2 - a * Decimal("1.5")
    return a, b, y100 - a - b



def parse_abaque(source) -> dict:
    """source : chemin ou fichier binaire. Lève ValueError si le fichier est invalide."""
    try:
        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"Fichier illisible : {e}") from e
    curves = _read_sheet(wb, "Abaque CPH", CURVE_COLUMNS)
    mappings = _read_sheet(wb, "Mappage inventaire", MAPPING_COLUMNS)
    valid_status = {s.value for s in CphCurve.Status}
    warnings = []
    for c in curves:
        if c["status"] not in valid_status:
            raise ValueError(f"{c['curve_id']} : statut inconnu {c['status']!r}")
        fa, fb, fc = fitted_coefficients(c["conso_50_l_h"], c["conso_75_l_h"], c["conso_100_l_h"])
        if max(abs(fa - c["coef_a"]), abs(fb - c["coef_b"]), abs(fc - c["coef_c"])) > Decimal("0.05"):
            warnings.append(f"{c['curve_id']} : coefficients abaque ({c['coef_a']}, {c['coef_b']}, {c['coef_c']}) "
                            f"≠ ajustement des 3 points ({fa:.3f}, {fb:.3f}, {fc:.3f})")
        if c["power_factor"] is None:
            warnings.append(f"{c['curve_id']} : cos φ absent — courbe inutilisable tant qu'il n'est pas renseigné.")
    known_ids = {c["curve_id"] for c in curves}
    for m in mappings:
        ids = [x.strip() for x in str(m["candidate_curve_ids"] or "").replace(";", ",").split(",") if x.strip()]
        unknown = [x for x in ids if x not in known_ids]
        if unknown:
            raise ValueError(f"Mappage « {m['inventory_label']} » : courbes inconnues {unknown}")
        m["candidate_curve_ids"] = ids
    return {
        "curves": curves, "mappings": mappings, "warnings": warnings,
        "status_counts": {s: sum(1 for c in curves if c["status"] == s) for s in valid_status},
    }


def apply_abaque(parsed: dict, file_name: str) -> list[str]:
    """Écrit courbes et mappages ; retourne les validations annulées (courbe plus candidate)."""
    reset = []
    with transaction.atomic():
        for c in parsed["curves"]:
            obj = CphCurve.objects.filter(curve_id=c["curve_id"]).first()
            fields = {k: v for k, v in c.items() if k != "curve_id"}
            if obj is None:
                CphCurve.objects.create(curve_id=c["curve_id"], abaque_file=file_name, **fields)
            else:
                fields.pop("power_factor")
                for k, v in fields.items():
                    setattr(obj, k, v)
                obj.abaque_file = file_name
                obj.imported_at = timezone.now()
                obj.save()
        for m in parsed["mappings"]:
            defaults = {k: v for k, v in m.items() if k not in ("inventory_label", "inventory_kva")}
            defaults.update(inventory_label_normalized=normalize_label(m["inventory_label"]),
                            abaque_file=file_name, imported_at=timezone.now())
            obj, _ = CphInventoryMapping.objects.update_or_create(
                inventory_label=m["inventory_label"], inventory_kva=m["inventory_kva"], defaults=defaults)
            if obj.validated_curve_id and obj.validated_curve.curve_id not in obj.candidate_curve_ids:
                if obj.match_status == CphInventoryMapping.MatchStatus.VALIDE_MANUELLEMENT:
                    reset.append(f"{obj.inventory_label} (courbe {obj.validated_curve.curve_id})")
                obj.validated_curve = None
                obj.validated_by = None
                obj.validated_at = None
                obj.validation_comment = ""
                obj.match_status = CphInventoryMapping.MatchStatus.A_VALIDER
                obj.match_method = ""
                obj.save()
        # Correspondance automatique des libellés fiables ; décisions humaines conservées.
        apply_auto_matching(CphInventoryMapping.objects.all(), list(CphCurve.objects.all()), timezone.now())
    return reset
