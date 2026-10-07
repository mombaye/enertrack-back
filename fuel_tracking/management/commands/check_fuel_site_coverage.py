"""
Vérifie (lecture seule) que tous les sites d'un fichier Ops apparaissent dans Suivi Carburant
(tableau « Suivis Consommations » / API cph/) : calculés (GE confirmé par Snowflake) ou listés
hors calcul avec leur motif. Liste les sites absents du référentiel et les écarts d'écriture.

    python manage.py check_fuel_site_coverage --file "Proposition_de_load_092026_v1.xlsx"

Le fichier doit contenir une colonne « Site_ID » (ou « Site ID ») dans l'une de ses feuilles.
Code de sortie 1 si au moins un site du fichier n'apparaît pas.
"""
import sys

import openpyxl
from django.core.management.base import BaseCommand, CommandError

HEADERS = {"site_id", "site id"}


def read_site_ids(path: str) -> tuple[str, list[str]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        col, ids = None, []
        for row in ws.iter_rows(max_row=ws.max_row, values_only=True):
            if col is None:
                for i, v in enumerate(row):
                    if isinstance(v, str) and v.strip().lower() in HEADERS:
                        col = i
                        break
                continue
            v = row[col] if col < len(row) else None
            if v not in (None, ""):
                ids.append(str(v).strip())
        if ids:
            return ws.title, ids
    raise CommandError("Aucune colonne Site_ID trouvée dans le fichier.")


class Command(BaseCommand):
    help = "Vérifie que tous les sites d'un fichier Ops apparaissent dans Suivi Carburant (lecture seule)."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True)

    def handle(self, *args, file, **opts):
        from core.models import Site
        from fuel_tracking.models import FuelConsommationMonthly, FuelSiteInventory

        sheet, ids = read_site_ids(file)
        unique = list(dict.fromkeys(ids))
        fcm = set(FuelConsommationMonthly.objects.values_list("site_id", flat=True).distinct())
        inv = dict(FuelSiteInventory.objects.values_list("site_id", "dg_count"))
        core = set(Site.objects.values_list("site_id", flat=True))
        known = fcm | set(inv)
        by_upper = {k.upper(): k for k in known | core}

        ge = [s for s in unique if (inv.get(s) or 0) > 0]
        shown = [s for s in unique if s in known]
        missing = [s for s in unique if s not in known]
        w = self.stdout.write
        w(f"Fichier : feuille « {sheet} », {len(ids)} ligne(s), {len(unique)} site(s) distinct(s)")
        w(f"  affichés dans Suivi Carburant          {len(shown)}")
        w(f"    calculés (GE confirmé, DG_COUNT > 0) {len(ge)}")
        w(f"    hors calcul (sans GE confirmé)       {len(shown) - len(ge)}")
        w(f"  absents du référentiel                 {len(missing)}")
        for s in missing:
            hint = by_upper.get(s.upper())
            if hint:
                w(self.style.WARNING(f"    {s} : écrit « {hint} » sur la plateforme (casse différente)"))
            elif s in core:
                w(self.style.WARNING(f"    {s} : présent dans core.Site mais ni dans l'inventaire Snowflake ni dans les imports Ops"))
            else:
                w(self.style.ERROR(f"    {s} : inconnu — lancer import_facturation_par_site ou la synchro Snowflake"))
        if missing:
            sys.exit(1)
        w(self.style.SUCCESS("OK — tous les sites du fichier apparaissent."))
