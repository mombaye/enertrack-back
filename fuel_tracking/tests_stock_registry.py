"""
Recette du registre unifié de stock carburant : moteur de cuve (calculateur de Steve, steve-v1), bibliothèque
de modèles, affectation site → cuve (Closing, refueling), conversion géométrique / calibrée, fichiers
Loliforce / SCSS, contrôles, gardiennage, comparaison, stock retenu (priorité), ENOC (classification
versionnée), historique refueling, pipeline du répertoire monté, file « À traiter », audit et API.
"""
import io
import math
import shutil
import tempfile
from datetime import datetime, timedelta
from datetime import timezone as dt_timezone
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from core.models import Site
from fuel_tracking.models import (FuelEnocMovement, FuelEnocTypeMapping, FuelGuardScope, FuelImportRow, FuelMovement,
                                  FuelRegistryCorrection, FuelStockObservation, FuelTankMaster, FuelTankModel)
from fuel_tracking.services import stock_enoc as E
from fuel_tracking.services import stock_import_files as F
from fuel_tracking.services import stock_import_pipeline as P
from fuel_tracking.services import stock_registry as R
from fuel_tracking.services import stock_tanks as T
from fuel_tracking.services import stock_work_queue as W
from fuel_tracking.services import tank_conversion as TC
from fuel_tracking.services import tank_geometry as G

HEAD = ["Site ID", "Site Name", "Catégorie", "Zone", "Ville", "MARQUE GE", "TYPE DE GE", "PUISSANCE GE", "Relevé en CM",
        "Relevé en litres", "Compteur horaire (h)", "Date de relevé", "Commentaires", "Site gardé par"]

# Valeurs de référence du calculateur de Steve (onglets Tank DB / Calculateur / Barèmes, Senegal_Tank_DB_Calculateur.xlsx).
STEVE_MODELS = [  # n°, famille, cote1, cote2, cote3, profil, Z, plaque, volume plein du fichier, k du fichier
    (1, "h-circ", 150, 0, 211, "ellip", 25, 4000, 4317.7264031668865, 0.9264134932371244),
    (2, "h-circ", 150, 0, 275, "ellip", 25, 5000, 5448.699758459212, 0.9176501223502732),
    (3, "h-circ", 124, 0, 162, "ellip", 25, 2400, 2358.9004485890646, 1.0174231818199528),
    (4, "rect", 77, 150, 200, "plat", 0, 2300, 2310.0, 0.9956709956709957),
    (5, "h-circ", 122, 0, 260, "ellip", 20, 3500, 3351.0949956236614, 1.044434730907599),
    (6, "h-circ", 150, 0, 160, "plat", 0, 3200, 2827.433388230814, 1.1317684842090334),
    (7, "h-circ", 200, 0, 330, "ellip", 25, 10000, 11414.45330784632, 0.8760822555668066),
    (8, "rect", 100, 150, 150, "plat", 0, 2250, 2250.0, 1.0),
]


def xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def guard_row(site, cm=None, l=None, d=datetime(2026, 10, 3), guard="LOLIFORCE", name=None, hm=1200):
    return [site, name or f"NAME_{site}", "A_Ax_S6", "DAKAR", "DAKAR", "SDMO", "K22", 22, cm, l, hm, d, None, guard]


def tank(site, tank_id="1", status="VALIDEE", family="rect", dims=(100, 100, 100), plate=1000, bottom="plat", bz=None, **kw):
    """Par défaut : parallélépipède 100 × 100 × 100 cm = 1 000 L, plaque 1 000 L (1 cm = 10 L)."""
    d1, d2, d3 = dims
    return FuelTankMaster.objects.create(
        site_id=site, tank_id=tank_id, validation_status=status, shape_family=family, shape=TC.FAMILY_TO_SHAPE.get(family),
        dimension_1_cm=d1, dimension_2_cm=d2, dimension_3_cm=d3, bottom_profile=bottom, bottom_z_cm=bz, top_profile="plat",
        capacity_plate_l=plate, **kw)


def spec(**kw):
    base = dict(site_id="S", tank_id="1", validated=True, status="VALIDEE")
    base.update(kw)
    return TC.TankSpec(**base)


class SteveEngineTests(TestCase):
    """Le moteur reproduit exactement le calculateur de Steve."""

    def test_full_volumes_and_k_match_calculator(self):
        for n, fam, d1, d2, d3, prof, z, plate, full_ref, k_ref in STEVE_MODELS:
            full = G.full_volume_l(fam, d1, d2, d3, prof, z)
            self.assertAlmostEqual(full, full_ref, places=6, msg=f"cuve {n}")
            self.assertAlmostEqual(G.calibration(plate, full)["k"], k_ref, places=9)

    def test_calculator_example_level_27cm(self):
        # Onglet Calculateur : cuve #2 (5000 L), niveau 0,27 m → géométrique 645,217 L, calibré 592,084 L.
        geo = G.volume_l("h-circ", 150, 0, 275, 27, "ellip", 25)
        self.assertAlmostEqual(geo, 645.2172586502544, places=6)
        self.assertAlmostEqual(geo * G.calibration(5000, G.full_volume_l("h-circ", 150, 0, 275, "ellip", 25))["k"], 592.0836963429139, places=6)

    def test_bareme_points(self):
        # Onglet Barèmes : 1 cm et 2 cm pour les cuves 1 à 8.
        ref = {1: [3.5265829276831235, 4.569605825179426, 2.4846170155645644, 15.4, 3.887425186705781, 2.6075572437407573, 6.291476514925616, 15],
               2: [10.027529053240507, 12.97170999829924, 7.086007520427092, 30.8, 11.029911640479089, 7.360452362646829, 17.859170345733244, 30]}
        for cm, values in ref.items():
            for (n, fam, d1, d2, d3, prof, z, *_), v in zip(STEVE_MODELS, values):
                self.assertAlmostEqual(G.volume_l(fam, d1, d2, d3, cm, prof, z), v, places=6, msg=f"cuve {n} à {cm} cm")

    def test_analytic_controls_from_notice(self):
        r = 50.0
        # Cylindre plein πR²L
        self.assertAlmostEqual(G.full_volume_l("h-circ", 100, 0, 200, "plat"), math.pi * r * r * 200 / 1000, places=3)
        # Deux fonds ellipsoïdaux : (4/3)πR²Z
        two_ends = G.full_volume_l("h-circ", 100, 0, 200, "ellip", 20) - math.pi * r * r * 200 / 1000
        self.assertAlmostEqual(two_ends, 4 / 3 * math.pi * r * r * 20 / 1000, places=2)
        # Deux hémisphères : (4/3)πR³ (Z = R imposé)
        hemi = G.full_volume_l("h-circ", 100, 0, 200, "hemi") - math.pi * r * r * 200 / 1000
        self.assertAlmostEqual(hemi, 4 / 3 * math.pi * r ** 3 / 1000, places=2)
        # Cuve verticale plate : πR²H ; tronconique : tronc de cône
        self.assertAlmostEqual(G.full_volume_l("v-cyl", 100, 150, 0), math.pi * r * r * 150 / 1000, places=3)
        self.assertAlmostEqual(G.full_volume_l("v-tronc", 100, 60, 120), math.pi * 120 / 3 * (50 ** 2 + 50 * 30 + 30 ** 2) / 1000, places=3)
        # Verticale à fond ellipsoïdal : fond = (2/3)πR²Z ; niveau = Z → exactement le fond
        self.assertAlmostEqual(G.volume_l("v-cyl", 100, 150, 0, 20, "ellip", 20), 2 / 3 * math.pi * r * r * 20 / 1000, places=2)
        # Klöpper : flèche 0,1938·D imposée
        self.assertAlmostEqual(G.end_rise_cm("klop", None, 150), 29.07, places=2)

    def test_geometry_errors(self):
        with self.assertRaises(G.GeometryError):
            G.full_volume_l("h-circ", None, 0, 200)
        with self.assertRaises(G.GeometryError):
            G.full_volume_l("h-circ", 100, 0, 200, "ellip", None)  # flèche Z manquante
        with self.assertRaises(G.GeometryError):
            G.full_volume_l("inconnue", 1, 1, 1)


class ConversionRuleTests(TestCase):
    """Règle de conversion cm → L : cuve VALIDÉE, volume calibré si plaque cohérente, sinon technique."""

    def test_calibrated_when_plate_coherent(self):
        t = spec(family="h-circ", d1=124, d3=162, bottom="ellip", bz=25, plate_capacity_l=2400)  # modèle 3 : k = 1,0174
        r = TC.convert(62, "CM", [t])
        self.assertEqual((r.status, r.method), ("VALID", "CALIBRATED"))
        self.assertAlmostEqual(r.geometric_l, 1179.450, places=2)
        self.assertAlmostEqual(r.calibrated_l, 1200.0, places=1)  # moitié de la plaque à mi-hauteur
        self.assertEqual(r.level_liters, r.calibrated_l)
        self.assertAlmostEqual(r.fill_pct, 50.0, places=1)

    def test_geometric_only_without_plate_or_incoherent_plate(self):
        r = TC.convert(27, "CM", [spec(family="h-circ", d1=150, d3=275, bottom="ellip", bz=25)])
        self.assertEqual((r.status, r.level_liters), ("GEOMETRIC_ONLY", None))
        self.assertAlmostEqual(r.geometric_l, 645.217, places=2)
        self.assertEqual(r.messages[0]["code"], "PLATE_MISSING")
        r = TC.convert(27, "CM", [spec(family="h-circ", d1=150, d3=275, bottom="ellip", bz=25, plate_capacity_l=5000)])  # écart 9 %
        self.assertEqual((r.status, r.level_liters, r.messages[0]["code"]), ("GEOMETRIC_ONLY", None, "PLATE_INCOHERENT"))
        self.assertAlmostEqual(r.calibrated_l, 592.084, places=2)  # calculé mais non retenu

    def test_no_generic_coefficient_without_validated_tank(self):
        r = TC.convert(20, "CM", [])
        self.assertEqual((r.level_liters, r.status, r.messages[0]["code"]), (None, "PENDING_TANK_MAPPING", "NO_VALIDATED_TANK"))
        for st, code in (("PROPOSEE", "TANK_PROPOSED"), ("A_CONFIRMER", "TANK_TO_CONFIRM"), ("INCOMPLETE", "TANK_INCOMPLETE")):
            r = TC.convert(20, "CM", [spec(family="rect", d1=100, d2=100, d3=100, plate_capacity_l=1000, validated=False, status=st)])
            self.assertEqual((r.level_liters, r.geometric_l, r.messages[0]["code"]), (None, None, code))

    def test_validated_curve_priority(self):
        t = spec(family="rect", d1=100, d2=100, d3=100, plate_capacity_l=1000,
                 calibration_curve=[[0, 0], [10, 50], [20, 300]], calibration_validated=True)
        r = TC.convert(15, "CM", [t])
        self.assertEqual((r.method, r.level_liters), ("TANK_CURVE", 175.0))

    def test_controls(self):
        t = spec(family="rect", d1=100, d2=100, d3=100, plate_capacity_l=1000)
        self.assertEqual(TC.convert(-3, "CM", [t]).messages[0]["code"], "NEGATIVE_LEVEL")
        self.assertEqual(TC.convert(120, "CM", [t]).messages[0]["code"], "LEVEL_ABOVE_HEIGHT")
        self.assertEqual(TC.convert(1100, "L", [t]).status, "INVALID")  # > 1000 L + 5 %
        self.assertEqual(TC.convert(1040, "L", [t]).status, "VALID")
        small = spec(family="rect", d1=100, d2=100, d3=100, plate_capacity_l=960)  # écart géométrie +4,2 % : cohérent
        self.assertEqual(TC.convert(100, "CM", [small]).status, "VALID")

    def test_multi_tank_without_id_never_tank_1(self):
        tanks = [spec(tank_id="1", family="rect", d1=100, d2=100, d3=100, plate_capacity_l=1000),
                 spec(tank_id="2", family="rect", d1=100, d2=100, d3=100, plate_capacity_l=1000)]
        r = TC.convert(20, "CM", tanks)
        self.assertEqual((r.level_liters, r.tank_id, r.messages[0]["code"]), (None, None, "MULTI_TANK_WITHOUT_ID"))
        self.assertEqual(TC.convert(20, "CM", tanks, tank_id="2").level_liters, 200.0)

    def test_native_liters(self):
        r = TC.convert(140, "L", [spec(family="rect", d1=100, d2=100, d3=100, plate_capacity_l=1000)])
        self.assertEqual((r.method, r.status, r.level_liters, r.messages), ("DIRECT_LITERS", "VALID", 140.0, []))
        self.assertEqual(TC.convert(140, "L", []).messages[0]["code"], "CAPACITY_NOT_VALIDATED")


class LevelCellTests(TestCase):
    def test_cells(self):
        p = F.parse_level_cell
        self.assertEqual([(r.value, r.unit, r.tank_id) for r in p(26, "CM").readings], [(26.0, "CM", None)])
        self.assertEqual([(r.value, r.unit) for r in p("16,5CM", "CM").readings], [(16.5, "CM")])
        self.assertEqual([(r.value, r.unit) for r in p("100 L", "L").readings], [(100.0, "L")])
        self.assertEqual([(r.value, r.tank_id) for r in p("Cuve 1=7 Cuve 2= 11", "CM").readings], [(7.0, "1"), (11.0, "2")])
        self.assertEqual([(r.value, r.tank_id) for r in p("C1: 8cm / C2: 24cm", "CM").readings], [(8.0, "1"), (24.0, "2")])
        multi = p("7 et 11", "CM")
        self.assertEqual([(r.value, r.tank_id) for r in multi.readings], [(7.0, None), (11.0, None)])
        self.assertEqual(p("Pas de jauge", "CM").error["code"], "MISSING_LEVEL")
        self.assertEqual(p("3 Bidons 20L", "CM").error["code"], "UNKNOWN_UNIT")
        self.assertEqual(p(None, "CM").readings, [])


class TmpRootMixin:
    def setUp(self):
        super().setUp()
        self.tmp = tempfile.mkdtemp()
        self.override = override_settings(FUEL_IMPORT_ROOT=self.tmp)
        self.override.enable()

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.tmp, ignore_errors=True)
        super().tearDown()


def steve_workbook() -> bytes:
    """Classeur au format du calculateur (onglets Tank DB + Profils) construit depuis les valeurs de référence."""
    head = ["#", "Country", "Type / Shape", "Capacity\n[Liters]", "Z\n[cm]", "D\n[cm]", "L\n[cm]", "General Mechanical Diagram",
            "Tank Mechanical Diagram", "Photos", "H\n[cm]  ← saisie", "Fuel Level\n[L]  ← résultat", "Famille", "Profil du fond",
            "Z saisi\n[cm]", "Cote 1\n[cm]", "Cote 2\n[cm]", "Cote 3\n[cm]", "Profil du ciel", "Zt saisi\n[cm]"]
    rows = [[], [], head]
    for n, fam, d1, d2, d3, prof, z, plate, *_ in STEVE_MODELS:
        rows.append([n, "Senegal", "Rectangular" if fam == "rect" else "Horizontal Dish Ends ", plate, z, d1, d3, None, None, None,
                     None, None, G.FAMILIES[fam], G.PROFILES[prof], z, d1, d2, d3, G.PROFILES["plat"], 0])
    rows.append([9, None, None])
    tb = G.tables()
    prof_rows = [["TABLES"], ["t / u", "Psi ellip", "Psi toris", "Psi korb", "Psi cone"]]
    for i in range(201):
        prof_rows.append([i * 0.005, tb["psi"]["ellip"][i], tb["psi"]["toris"][i], tb["psi"]["korb"][i], tb["psi"]["cone"][i]])
    return xlsx({"Tank DB": rows, "Profils": prof_rows})


class TankLibraryAndAssignmentTests(TmpRootMixin, TestCase):
    def setUp(self):
        super().setUp()
        for sid in ("DKR_0001", "DKR_0002", "DKR_0003", "DKR_0004"):
            Site.objects.create(site_id=sid, name=sid)

    def test_import_steve_models(self):
        b = P.process_bytes(steve_workbook(), "Senegal_Tank_DB_Calculateur.xlsx", "TANK_MODELS", origin="UPLOAD")
        self.assertEqual(FuelTankModel.objects.count(), 8)
        st = dict(FuelTankModel.objects.values_list("tank_model_id", "model_status"))
        self.assertEqual(st["SEN-TDB-03"], "COMPLET")
        self.assertEqual(st["SEN-TDB-02"], "A_CONFIRMER")  # écart +9,0 %
        self.assertEqual(st["SEN-TDB-07"], "A_CONFIRMER")  # écart +14,1 %
        self.assertEqual(sorted(st.values()).count("COMPLET"), 4)
        m2 = FuelTankModel.objects.get(tank_model_id="SEN-TDB-02")
        self.assertAlmostEqual(float(m2.geometric_full_l), 5448.700, places=2)
        self.assertEqual((m2.source_sheet, m2.shape_code, m2.bottom_profile), ("Tank DB", "h-circ", "ellip"))
        self.assertEqual(b.summary["profils"]["statut"], "IDENTIQUES")
        self.assertEqual(FuelTankMaster.objects.count(), 0)  # un modèle n'est jamais une cuve de site
        again = P.process_bytes(steve_workbook() + b"", "copie.xlsx", "TANK_MODELS", origin="UPLOAD")
        self.assertEqual(again.status, "REJECTED")  # même empreinte : doublon technique

    def closing(self, rows):
        head = ["x", "ID du site", "Nom du site", "Forme Géométrique Cuve", "Longueur [m]", "Largeur [m]", "Hauteur [m]", "Diamètre [m]", "Capacité cuve (l)"]
        return P.process_bytes(xlsx({"Synthese conso fuel": [[], head] + rows}), "closing.xlsx", "CLOSING", origin="UPLOAD")

    def test_closing_assignment_statuses_and_model_proposal(self):
        P.process_bytes(steve_workbook(), "steve.xlsx", "TANK_MODELS", origin="UPLOAD")
        geo = math.pi * 0.62 ** 2 * 1.62 * 1000  # 1956 L, fonds plats (calcul du Closing)
        self.closing([
            [None, "DKR_0001", "A", "Cuve Cylindrique Horizontale", 1.62, None, None, 1.24, geo],          # cohérent → VALIDÉE + modèle 3 proposé
            [None, "DKR_0002", "B", "Cuve Cylindrique Horizontale", 2.7, None, None, 1.5, None],           # cotes sans capacité → PROPOSÉE
            [None, "DKR_0003", "C", "Cuve Parallélépipédique", 2.0, None, 0.77, None, 1500],               # largeur manquante, capacité → À_CONFIRMER
            [None, "DKR_0004", "D", None, None, None, None, None, None],                                  # aucune info
        ])
        t1 = FuelTankMaster.objects.get(site_id="DKR_0001")
        self.assertEqual((t1.validation_status, t1.model_match_status, t1.tank_model.tank_model_id), ("VALIDEE", "PROPOSED", "SEN-TDB-03"))
        self.assertIsNone(t1.capacity_plate_l)  # capacité recalculée par le Closing ≠ plaque ; plaque du modèle non utilisée (proposition)
        self.assertEqual((float(t1.dimension_1_cm), float(t1.dimension_3_cm), t1.bottom_profile), (124.0, 162.0, "plat"))
        self.assertIn("pas une capacité de plaque", t1.validation_reason)  # capacité Closing recalculée : jamais présentée comme plaque
        self.assertEqual(FuelTankMaster.objects.get(site_id="DKR_0002").validation_status, "PROPOSEE")
        self.assertEqual(FuelTankMaster.objects.get(site_id="DKR_0003").validation_status, "A_CONFIRMER")
        self.assertFalse(FuelTankMaster.objects.filter(site_id="DKR_0004").exists())

    def test_closing_takes_over_refueling_proposal_and_survives_evaluation(self):
        geo = math.pi * 0.62 ** 2 * 1.62 * 1000
        t = tank("DKR_0001", status="PROPOSEE", family="h-circ", dims=(124, None, 162), plate=None, source_type="HISTORICAL_REFUELING")
        self.closing([[None, "DKR_0001", "A", "Cuve Cylindrique Horizontale", 1.62, None, None, 1.24, geo]])
        t.refresh_from_db()
        self.assertEqual((t.source_type, t.validation_status), ("CLOSING", "VALIDEE"))
        T.evaluate_site_tanks()
        t.refresh_from_db()
        self.assertEqual(t.validation_status, "VALIDEE")  # la réévaluation applique la règle Closing, pas celle du refueling

    def test_evaluation_repairs_tank_updated_by_closing_with_stale_source(self):
        geo = math.pi * 0.62 ** 2 * 1.62 * 1000
        b = self.closing([[None, "DKR_0001", "A", "Cuve Cylindrique Horizontale", 1.62, None, None, 1.24, geo]])
        FuelTankMaster.objects.filter(site_id="DKR_0001").update(source_type="HISTORICAL_REFUELING", validation_status="PROPOSEE")
        r = T.evaluate_site_tanks()
        t = FuelTankMaster.objects.get(site_id="DKR_0001")
        self.assertEqual((t.source_batch_id, t.source_type, t.validation_status, r["sources_corrigees"]), (b.pk, "CLOSING", "VALIDEE", 1))

    def test_model_never_proposed_on_capacity_alone_or_ambiguity(self):
        P.process_bytes(steve_workbook(), "steve.xlsx", "TANK_MODELS", origin="UPLOAD")
        models = T.library_models()
        m, why = T.propose_model(spec(family="h-circ", d1=160, d3=260), models)  # 5000 L courant mais cotes différentes
        self.assertIsNone(m)
        self.assertIn("Aucun modèle", why)
        m, _ = T.propose_model(spec(family="h-circ", d1=150, d3=273), models)  # cote L à 2 cm du modèle 2
        self.assertEqual(m.tank_model_id, "SEN-TDB-02")
        m, _ = T.propose_model(spec(family=None), models)
        self.assertIsNone(m)

    def test_confirming_model_applies_geometry_and_plate(self):
        P.process_bytes(steve_workbook(), "steve.xlsx", "TANK_MODELS", origin="UPLOAD")
        t = tank("DKR_0001", family="h-circ", dims=(124, None, 162), plate=None, status="VALIDEE")
        t.tank_model = FuelTankModel.objects.get(tank_model_id="SEN-TDB-03")
        t.model_match_status = "PROPOSED"
        t.save()
        R.correct_tank(t, {"model_match_status": "CONFIRMED"}, None, "Plaque lue sur site")
        t.refresh_from_db()
        self.assertEqual((t.bottom_profile, float(t.bottom_z_cm), float(t.capacity_plate_l)), ("ellip", 25.0, 2400.0))
        r = TC.convert(62, "CM", [TC.TankSpec.from_model(t)])
        self.assertEqual((r.method, round(r.level_liters)), ("CALIBRATED", 1200))

    def test_refueling_plate_evidence(self):
        tank("DKR_0001", family="h-circ", dims=(150, None, 280), plate=None, status="VALIDEE", source_type="CLOSING")  # 4948 L
        tank("DKR_0002", family="h-circ", dims=(150, None, 200), plate=None, status="VALIDEE", source_type="CLOSING")  # 3534 L
        b = P.process_bytes(xlsx({"x": [["x"]]}), "vide.xlsx", "HISTORICAL_REFUELING", origin="UPLOAD")
        stats = {}
        T.apply_plate_evidence(b, {"DKR_0001": [(5000.0, "ligne 3")], "DKR_0002": [(5000.0, "ligne 4")],
                                   "DKR_0003": [(1000.0, "ligne 5")]}, stats)
        self.assertEqual(float(FuelTankMaster.objects.get(site_id="DKR_0001").capacity_plate_l), 5000.0)  # écart -1,0 %
        self.assertIsNone(FuelTankMaster.objects.get(site_id="DKR_0002").capacity_plate_l)  # écart -29 % : refusé
        t3 = FuelTankMaster.objects.get(site_id="DKR_0003")
        self.assertEqual((t3.validation_status, float(t3.capacity_plate_l)), ("A_CONFIRMER", 1000.0))
        self.assertEqual((stats["plaques_appliquees"], stats["plaques_incoherentes"], stats["cuves_capacite_seule"]), (1, 1, 1))
        self.assertIsNone(T.plate_from_cuve_cell("4418"))  # valeur recalculée, pas une plaque
        self.assertIsNone(T.plate_from_cuve_cell("Bidon"))
        self.assertEqual(T.plate_from_cuve_cell(5000), 5000.0)


class GuardFileImportTests(TmpRootMixin, TestCase):
    """Fichiers Loliforce / SCSS : contrôles ligne à ligne et validation."""

    def setUp(self):
        super().setUp()
        for sid, name in (("DKR_0001", "ALPHA"), ("DKR_0002", "BETA"), ("DKR_0003", "GAMMA"), ("DKR_0004", "DELTA")):
            Site.objects.create(site_id=sid, name=name)

    def run_file(self, rows, source="LOLIFORCE", sheet="KOLDA", name="f.xlsx"):
        return P.process_bytes(xlsx({sheet: [["SEMAINE 40"], [], HEAD] + rows}), name, source, origin="UPLOAD")

    def scope(self, site, provider="LOLIFORCE"):
        FuelGuardScope.objects.create(site_id=site, provider=provider, active_from="2026-01-01")

    def test_controls_unknown_site_missing_date_missing_level_unknown_unit(self):
        b = self.run_file([guard_row("XXX_9999", cm=10), guard_row("DKR_0001", cm=10, d=None),
                           guard_row("DKR_0002", cm="Pas de jauge"), guard_row("DKR_0003", cm="3 Bidons 20L")])
        codes = {r.row_number: [m["code"] for m in r.messages] for r in b.rows.all()}
        self.assertIn("UNKNOWN_SITE", codes[4])
        self.assertIn("MISSING_DATE", codes[5])
        self.assertIn("MISSING_LEVEL", codes[6])
        self.assertIn("UNKNOWN_UNIT", codes[7])
        self.assertEqual((b.rows_rejected, b.status, FuelStockObservation.objects.count()), (4, "REJECTED", 0))

    def test_cm_without_validated_tank_has_explicit_action(self):
        self.run_file([guard_row("DKR_0001", cm=18)])
        o = FuelStockObservation.objects.get()
        self.assertEqual((o.level_liters, o.conversion_status, o.validation_status), (None, "PENDING_TANK_MAPPING", "PENDING"))
        self.assertEqual(o.issue_codes[0], "NO_VALIDATED_TANK")
        self.assertIn("Renseigner ou confirmer la cuve avant conversion", o.validation_reason)
        self.assertEqual(R.display_status(o), "A_TRAITER")

    def test_full_chain_auto_validated_calibrated(self):
        tank("DKR_0001")
        self.scope("DKR_0001")
        b = self.run_file([guard_row("DKR_0001", cm=18)])
        o = FuelStockObservation.objects.get()
        self.assertEqual((o.validation_status, o.conversion_method, float(o.level_liters), float(o.volume_geometric_l)),
                         ("VALIDATED", "CALIBRATED", 180.0, 180.0))
        self.assertEqual((R.display_status(o), o.issue_codes, b.status), ("AUTO_VALIDE", [], "ACCEPTED"))
        self.assertIn("Volume calibré", o.conversion_rule)

    def test_geometric_only_is_technical_until_explicit_manual_validation(self):
        tank("DKR_0001", plate=None)
        self.scope("DKR_0001")
        self.run_file([guard_row("DKR_0001", cm=18)])
        o = FuelStockObservation.objects.get()
        self.assertEqual((o.level_liters, float(o.volume_geometric_l), o.issue_codes), (None, 180.0, ["PLATE_MISSING"]))
        self.assertIsNone(R.fuel_stock_retained()["DKR_0001"]["stock_l"])
        R.validate_observation(o, get_user_model().objects.create_user("m", "m@x.sn", "p", role="manager"), "Plaque illisible, volume géométrique accepté")
        o.refresh_from_db()
        self.assertEqual((float(o.level_liters), R.display_status(o)), (180.0, "VALIDE_MANUELLEMENT"))
        self.assertEqual(R.fuel_stock_retained()["DKR_0001"]["stock_l"], 180.0)

    def test_negative_and_over_capacity_rejected_but_kept(self):
        tank("DKR_0001", plate=500, dims=(100, 100, 50))
        self.scope("DKR_0001")
        self.run_file([guard_row("DKR_0001", cm=-5), guard_row("DKR_0001", l=900, d=datetime(2026, 10, 4))])
        obs = {o.level_raw: o for o in FuelStockObservation.objects.all()}
        self.assertEqual(obs["-5"].validation_status, "REJECTED")
        self.assertEqual((obs["900"].conversion_status, obs["900"].validation_status, obs["900"].issue_codes), ("INVALID", "REJECTED", ["OVER_CAPACITY"]))

    def test_scope_not_configured_and_provider_mismatch_never_rejected(self):
        tank("DKR_0001")
        tank("DKR_0002")
        self.scope("DKR_0002", "SCSS")
        self.run_file([guard_row("DKR_0001", cm=18), guard_row("DKR_0002", cm=18)])
        o1 = FuelStockObservation.objects.get(site_id="DKR_0001")
        o2 = FuelStockObservation.objects.get(site_id="DKR_0002")
        self.assertEqual((o1.guard_scope_status, o1.validation_status, o1.issue_codes), ("NOT_CONFIGURED", "PENDING", ["NOT_CONFIGURED"]))
        self.assertEqual((o2.guard_scope_status, o2.validation_status), ("PROVIDER_MISMATCH", "PENDING"))
        self.assertEqual(float(o2.level_liters), 180.0)

    def test_cm_and_liters_both_kept_with_gap(self):
        tank("DKR_0001")
        self.scope("DKR_0001")
        self.run_file([guard_row("DKR_0001", cm=10, l=105), guard_row("DKR_0001", cm=10, l=400, d=datetime(2026, 10, 4))])
        ok = FuelStockObservation.objects.get(level_raw="105")
        self.assertEqual((float(ok.declared_liters), float(ok.level_cm), float(ok.cm_liters_gap_l), ok.validation_status),
                         (105.0, 10.0, 5.0, "VALIDATED"))
        gap = FuelStockObservation.objects.get(level_raw="400")
        self.assertEqual((float(gap.cm_liters_gap_l), gap.validation_status, gap.issue_codes), (300.0, "PENDING", ["CM_L_GAP"]))
        self.assertEqual(float(gap.declared_liters), 400.0)  # valeur source jamais écrasée

    def test_old_reading_flagged_but_kept(self):
        tank("DKR_0001")
        self.scope("DKR_0001")
        b = self.run_file([guard_row("DKR_0001", cm=10, d=datetime(2026, 1, 5))])
        self.assertIn("OLD_READING", [m["code"] for m in b.rows.get().messages])
        self.assertEqual(FuelStockObservation.objects.get().validation_status, "VALIDATED")

    def test_multi_tank_cell_without_id(self):
        tank("DKR_0001", tank_id="1")
        tank("DKR_0001", tank_id="2")
        self.run_file([guard_row("DKR_0001", cm="7 et 11"), guard_row("DKR_0001", cm="C1=7 / C2=11", d=datetime(2026, 10, 4))])
        without = FuelStockObservation.objects.filter(observation_datetime__day=3)
        self.assertTrue(all(o.tank_id is None and o.level_liters is None and "MULTI_TANK_WITHOUT_ID" in o.issue_codes for o in without))
        labelled = FuelStockObservation.objects.filter(observation_datetime__day=4).order_by("tank_id")
        self.assertEqual([(o.tank_id, float(o.level_liters)) for o in labelled], [("1", 70.0), ("2", 110.0)])

    def test_duplicates_row_and_file(self):
        rows = [guard_row("DKR_0001", cm=18), guard_row("DKR_0001", cm=18)]
        b1 = self.run_file(rows)
        self.assertEqual((FuelStockObservation.objects.count(), b1.rows_duplicate), (1, 1))
        same = self.run_file(rows)
        self.assertEqual((same.status, same.duplicate_of_id), ("REJECTED", b1.pk))
        b3 = self.run_file(rows + [guard_row("DKR_0001", cm=15, d=datetime(2026, 10, 10))], name="f2.xlsx")
        self.assertEqual((FuelStockObservation.objects.count(), b3.rows_duplicate), (2, 2))

    def test_same_source_conflicting_value_goes_to_review(self):
        tank("DKR_0001")
        self.scope("DKR_0001")
        self.run_file([guard_row("DKR_0001", cm=18), guard_row("DKR_0001", cm=40)])
        obs = sorted(FuelStockObservation.objects.all(), key=lambda o: o.id)
        self.assertEqual([o.validation_status for o in obs], ["VALIDATED", "PENDING"])
        self.assertEqual(obs[1].issue_codes, ["SAME_SOURCE_CONFLICT"])

    def test_block_without_site_id_resolved_by_name_and_compile_sheet_skipped(self):
        rows = [["SEMAINE 41"], [], HEAD[1:], guard_row("x", cm=12, name="BETA")[1:]]
        b = P.process_bytes(xlsx({"KAOLACK": rows, "Compile": [HEAD, guard_row("DKR_0003", cm=5)]}), "g.xlsx", "LOLIFORCE", origin="UPLOAD")
        self.assertEqual(FuelStockObservation.objects.get().site_id, "DKR_0002")
        self.assertIn("Compile", b.summary["feuilles_ignorees"])


class ComparisonAndRetainedTests(TestCase):
    def setUp(self):
        Site.objects.create(site_id="DKR_0001", name="ALPHA")
        self.t = tank("DKR_0001", dims=(100, 100, 200), plate=2000)
        self.now = datetime(2026, 10, 14, 12, tzinfo=dt_timezone.utc)

    def obs(self, source, liters, hours_ago, status="VALIDATED", tank_id="1", proof=None, site="DKR_0001"):
        return FuelStockObservation.objects.create(
            source_type=source, source_event_id=f"{source}-{liters}-{hours_ago}-{tank_id}-{proof}", site_id=site, tank_id=tank_id,
            observation_datetime=self.now - timedelta(hours=hours_ago), level_raw=str(liters), level_value=liters, level_unit="L",
            level_liters=liters, conversion_method="DIRECT_LITERS", conversion_status="VALID", validation_status=status, proof_reference=proof)

    def test_coherent_within_tolerance(self):
        self.obs("LOLIFORCE", 500, 10)
        self.obs("SCSS", 680, 30)  # écart 180 ≤ max(100, 200)
        c = R.guard_comparisons("DKR_0001")
        self.assertEqual((c[0]["status"], c[0]["tolerance_l"]), (R.COHERENT, 200.0))

    def test_conflict_both_to_review_no_average(self):
        a = self.obs("LOLIFORCE", 500, 10)
        b = self.obs("SCSS", 900, 30)
        self.assertEqual(R.apply_guard_conflicts(["DKR_0001"]), 2)
        a.refresh_from_db()
        b.refresh_from_db()
        self.assertEqual((a.validation_status, b.validation_status, a.issue_codes), ("PENDING", "PENDING", ["STOCK_CONFLICT"]))
        self.assertIsNone(R.fuel_stock_retained(now=self.now)["DKR_0001"]["stock_l"])

    @override_settings(FUEL_GUARD_CONFLICT_PRIORITY="SCSS")
    def test_conflict_with_configured_priority(self):
        a = self.obs("LOLIFORCE", 500, 10)
        b = self.obs("SCSS", 900, 30)
        R.apply_guard_conflicts(["DKR_0001"])
        a.refresh_from_db()
        b.refresh_from_db()
        self.assertEqual((a.validation_status, b.validation_status), ("PENDING", "VALIDATED"))
        self.assertEqual(R.fuel_stock_retained(now=self.now)["DKR_0001"]["stock_l"], 900.0)

    def test_single_tank_site_compares_without_tank_id(self):
        self.obs("LOLIFORCE", 500, 10, tank_id=None)
        self.obs("SCSS", 900, 30, tank_id="1")
        self.assertEqual(R.guard_comparisons("DKR_0001")[0]["status"], R.STOCK_CONFLICT)

    def test_not_comparable_beyond_48h_or_without_capacity(self):
        self.obs("LOLIFORCE", 500, 10)
        self.obs("SCSS", 900, 70)
        self.assertEqual(R.guard_comparisons("DKR_0001"), [])
        self.assertEqual(R.compare_pair(*FuelStockObservation.objects.all()[:2], None)["status"], R.NOT_COMPARABLE)

    def test_human_validated_never_downgraded(self):
        a = self.obs("LOLIFORCE", 500, 10)
        a.validated_by = get_user_model().objects.create_user("m", "m@x.sn", "p", role="manager")
        a.save()
        self.obs("SCSS", 900, 30)
        R.apply_guard_conflicts(["DKR_0001"])
        a.refresh_from_db()
        self.assertEqual(a.validation_status, "VALIDATED")

    def test_retained_requires_validated_tank(self):
        self.t.validation_status = "PROPOSEE"
        self.t.save()
        self.obs("LOLIFORCE", 600, 20)
        r = R.fuel_stock_retained(now=self.now)["DKR_0001"]
        self.assertIsNone(r["stock_l"])
        self.assertIn("Aucune cuve validée", r["unavailable_reason"])

    def test_retained_latest_day_then_priority(self):
        self.obs("ENOC", 700, 5, status="PENDING")
        self.obs("LOLIFORCE", 2500, 8)  # au-delà de la capacité : jamais retenu
        self.obs("ENOC", 610, 9)
        self.obs("SCSS", 600, 11)
        self.obs("SNOWFLAKE", 999, 1)
        r = R.fuel_stock_retained(now=self.now)["DKR_0001"]
        self.assertEqual((r["stock_l"], r["source_type"], r["statut_releve"]), (600.0, "SCSS", "AUTO_VALIDE"))  # gardiennage > ENOC le même jour
        self.obs("ENOC", 620, 10, proof="photo.jpg")
        self.assertEqual(R.fuel_stock_retained(now=self.now)["DKR_0001"]["stock_l"], 620.0)  # terrain avec preuve en premier
        self.obs("ENOC", 400, 30)  # jour précédent : jamais préféré au jour le plus récent
        self.assertEqual(R.fuel_stock_retained(now=self.now)["DKR_0001"]["stock_l"], 620.0)

    def test_retained_blocked_when_reading_mentions_unknown_tank(self):
        self.obs("LOLIFORCE", 600, 2, tank_id="1")
        FuelStockObservation.objects.create(
            source_type="LOLIFORCE", source_event_id="c2", site_id="DKR_0001", tank_id="2", observation_datetime=self.now - timedelta(hours=2),
            level_raw="Cuve 2: 10", level_value=10, level_unit="CM", conversion_status="PENDING_TANK_MAPPING", validation_status="PENDING",
            issue_codes=["TANK_UNKNOWN"])
        r = R.fuel_stock_retained(now=self.now)["DKR_0001"]
        self.assertIsNone(r["stock_l"])
        self.assertIn("cuve non référencée", r["unavailable_reason"])

    def test_retained_multi_tank_sum_and_incomplete(self):
        tank("DKR_0001", tank_id="2", plate=1000)
        self.obs("LOLIFORCE", 600, 2, tank_id="1")
        r = R.fuel_stock_retained(now=self.now)["DKR_0001"]
        self.assertIsNone(r["stock_l"])
        self.assertIn("cuve(s) 2", r["unavailable_reason"])
        self.obs("LOLIFORCE", 300, 3, tank_id="2")
        r = R.fuel_stock_retained(now=self.now)["DKR_0001"]
        self.assertEqual((r["stock_l"], r["capacity_l"], r["tanks_count"]), (900.0, 3000.0, 2))


class EnocRegistryTests(TestCase):
    """ENOC : réutilisation de FuelEnocMovement (aucune nouvelle connexion), classification versionnée."""

    def setUp(self):
        Site.objects.create(site_id="DKR_0001", name="ALPHA")
        tank("DKR_0001")
        self.when = timezone.now() - timedelta(days=1)

    def mk(self, code, status="done", operation=True, op_type="AJOUT", **raw_extra):
        op = {"request_code": code, "quantity_added_liters": 150, "operation_type": op_type, **raw_extra} if operation else None
        raw = {"source_id": code, "status": status, "site_id": "dkr_0001", "quantity_added_liters": op and op["quantity_added_liters"],
               "request": {"request_code": code}, "operation": op}
        return FuelEnocMovement.objects.create(
            source_id=code, request_code=code, status=status, site_id="dkr_0001", operation_type=op_type, operation_date=self.when,
            created_at_source=self.when - timedelta(hours=2), level_before=Decimal("20"), level_before_unit="cm",
            level_after=Decimal("35"), level_after_unit="cm" if operation else None, quantity_added_liters=Decimal("0"),
            technician_name="Tech A", raw_payload=raw)

    def test_unknown_type_never_a_real_movement(self):
        self.mk("REQ-1")
        E.sync_enoc_to_registry()
        mv = FuelMovement.objects.get()
        self.assertEqual((mv.source_category, mv.movement_type, mv.retained_in_stock, mv.raw_operation_type), ("INCONNU", None, False, "AJOUT"))
        self.assertIn("ENOC_TYPE_UNKNOWN", [a["code"] for a in mv.alerts])
        item = W.build("mouvements_enoc_inconnus")[0]
        self.assertEqual((item["action_type"], item["ref"]["raw_value"]), ("MAP_ENOC", "AJOUT"))

    def test_configured_mapping_versioned_and_reclassifies(self):
        self.mk("REQ-1", supplier="EYDON", photos=["https://x/p1.jpg"], ticket="T42", op_type="LIVRAISON")
        E.sync_enoc_to_registry()
        E.set_enoc_mapping("livraison", "AJOUT", None, "première version")
        E.set_enoc_mapping("LIVRAISON", "LIVRAISON_CONFIRMEE", None, "validé avec Ops")
        self.assertEqual(list(FuelEnocTypeMapping.objects.filter(raw_value="LIVRAISON").values_list("version", "active")),
                         [(2, True), (1, False)])
        E.sync_enoc_to_registry()
        mv = FuelMovement.objects.get(source_event_id="REQ-1")
        self.assertEqual((mv.source_category, mv.movement_type, mv.supplier, float(mv.quantity_l), mv.retained_in_stock),
                         ("LIVRAISON_CONFIRMEE", "DELIVERY_EYDON", "EYDON", 150.0, True))
        self.assertIn("version 2", mv.mapping_rule)
        self.assertTrue(FuelRegistryCorrection.objects.filter(entity="MOVEMENT", field="source_category").exists())
        with self.assertRaises(ValueError):
            E.set_enoc_mapping("LIVRAISON", "AJOUT", None, "")  # motif obligatoire

    def test_confirmed_delivery_without_proof_not_retained(self):
        E.set_enoc_mapping("LIVRAISON", "LIVRAISON_CONFIRMEE", None, "ok")
        self.mk("REQ-2", op_type="LIVRAISON")
        E.sync_enoc_to_registry()
        mv = FuelMovement.objects.get()
        self.assertFalse(mv.retained_in_stock)
        self.assertIn("DELIVERY_WITHOUT_PROOF", [a["code"] for a in mv.alerts])

    def test_request_only_and_total_card(self):
        E.set_enoc_mapping("CARTE", "CARTE_TOTAL", None, "ok")
        self.mk("REQ-3", status="pending", operation=False)
        self.mk("REQ-4", op_type="CARTE")
        E.sync_enoc_to_registry()
        req = FuelMovement.objects.get(source_event_id="REQ-3")
        self.assertEqual((req.source_category, req.quantity_l, req.retained_in_stock), ("DEMANDE_LIVRAISON", None, False))
        card = FuelMovement.objects.get(source_event_id="REQ-4")
        self.assertEqual((card.source_category, card.movement_type, card.retained_in_stock), ("CARTE_TOTAL", "DELIVERY_TOTAL_CARD", False))
        queued = {i["ref"].get("movement_id") for i in W.build("mouvements_sans_preuve")}
        self.assertFalse(queued & {req.id, card.id})  # ni une demande ni une carte Total ne sont des livraisons sans preuve
        with self.assertRaises(ValueError):
            card.quantity_l = Decimal("150")
            R.validate_movement(card, None, "test")

    def test_observations_and_idempotence(self):
        self.mk("REQ-5")
        E.sync_enoc_to_registry()
        obs = {o.source_event_id: o for o in FuelStockObservation.objects.filter(source_type="ENOC")}
        self.assertEqual(float(obs["REQ-5:before"].level_liters), 200.0)
        self.assertEqual(obs["REQ-5:after"].validation_status, "VALIDATED")
        s2 = E.sync_enoc_to_registry()
        self.assertEqual((s2["mouvements_crees"], s2["releves_crees"], s2["inchanges"], s2["lot"]), (0, 0, 1, None))

    def test_enoc_dictionary(self):
        self.mk("REQ-6", op_type="Ajout ")
        d = E.enoc_type_dictionary()
        self.assertEqual((d[0]["raw_value"], d[0]["categorie"], d[0]["suggestion"], d[0]["configure"]), ("AJOUT", "INCONNU", "AJOUT", False))


class RefuelingHistoryTests(TmpRootMixin, TestCase):
    def setUp(self):
        super().setUp()
        Site.objects.create(site_id="DKR_0061", name="PEAGE10")
        Site.objects.create(site_id="THS_0084", name="BAITY")

    def test_history_movements_and_classification(self):
        head = ["N° Ticket FMS", "Site ID", "SITE NAME", "ZONE", "VILLE", "Typologie", "retrait", "Source", "Date derniére refueling",
                "Date de la demande", "Responsable", "Source Site ID/Dépôt- ID Destination", "QTT initiale mesurée sur site (L)",
                "QTT initiale mesurée sur site (cm)", "Quantité demandée (L)", "Qtt finale mesurée sur sites (L)",
                "Qtt finale mesurée sur sites (cm)", "Compteur Horaire derniére refueling (H)", "Compteur Horaire Actuel (H)",
                "Delta CH (H)", "CUVE ", "Qtt initiale RMS", "Qtt finale RMS", "RMS DG RH ", "Motif Transfert", "Validé par", "Statut", "Commentaire"]
        d = datetime(2026, 9, 1)
        rows = [
            [5760795, "DKR_0061", "P", "DAKAR", "R", "A", "Ajout", "TOTAL", None, d, "Waly", None, 217, 13, 100, 354, 18, 6090, 6105, 15, 5000, None, None, None, "Appoint", None, "Soldé", None],
            [5760796, "THS_0084", "B", "THIES", "T", "A", "Ajout", "EYDON", None, d, "Malick", None, 0, None, 1000, None, None, None, None, 0, "Bidon", None, None, None, "Dépotage", None, "En cours", None],
            [5760797, "THS_0084", "B", "THIES", "T", "A", "retrait", "Ponction", None, d, "Malick", "DKR_0061", 300, None, 100, 200, None, None, None, 0, 1000, None, None, None, "Ponction", None, "Soldé", None],
        ]
        b = P.process_bytes(xlsx({"Suivi refueling septembre": [[None], head] + rows}), "r.xlsx", "HISTORICAL_REFUELING", origin="UPLOAD")
        mvs = {m.ticket_reference: m for m in FuelMovement.objects.all()}
        self.assertEqual((mvs["5760795"].movement_type, mvs["5760795"].source_category), ("DELIVERY_TOTAL_CARD", "CARTE_TOTAL"))
        self.assertEqual(mvs["5760796"].source_category, "DEMANDE_LIVRAISON")  # Eydon non soldé : demande seulement
        self.assertEqual((mvs["5760797"].movement_type, mvs["5760797"].counterpart_site_id), ("TRANSFER_OUT", "DKR_0061"))
        self.assertTrue(all(m.quantity_l is None and not m.retained_in_stock for m in mvs.values()))
        self.assertEqual(FuelStockObservation.objects.count(), 0)
        self.assertEqual(b.summary["mouvements_crees"], 3)
        # CUVE 5000 / 1000 → capacité de plaque déclarée (cuve créée À_CONFIRMER : aucune cote connue)
        t = FuelTankMaster.objects.get(site_id="DKR_0061")
        self.assertEqual((t.validation_status, float(t.capacity_plate_l)), ("A_CONFIRMER", 5000.0))


class PipelineTests(TmpRootMixin, TestCase):
    def setUp(self):
        super().setUp()
        Site.objects.create(site_id="DKR_0001", name="ALPHA")

    def test_mounted_directory_flow(self):
        root = Path(self.tmp)
        P.ensure_dirs()
        for sub in ("incoming/loliforce", "incoming/scss", "incoming/closing", "incoming/modeles_cuves", "processing", "accepted",
                    "rejected", "archive"):
            self.assertTrue((root / sub).is_dir())
        self.assertFalse((root / "incoming/eydon").exists())
        (root / "incoming/modeles_cuves/Senegal_Tank_DB_Calculateur.xlsx").write_bytes(steve_workbook())
        data = xlsx({"KOLDA": [HEAD, guard_row("DKR_0001", cm=12), guard_row("XXX_1", cm=12)]})
        src = root / "incoming/loliforce/semaine41.xlsx"
        src.write_bytes(data)
        models_batch, b = P.scan_incoming()  # modèles traités avant les relevés
        self.assertEqual((models_batch.source_type, FuelTankModel.objects.count()), ("TANK_MODELS", 8))
        self.assertFalse(src.exists())
        self.assertEqual(b.status, "VALIDATION_REQUIRED")
        self.assertEqual(Path(b.stored_path).read_bytes(), data)
        self.assertIn("UNKNOWN_SITE", Path(b.report_path).read_text(encoding="utf-8-sig"))
        src.write_bytes(data)
        (dup,) = P.scan_incoming()
        self.assertEqual((dup.status, dup.duplicate_of_id), ("REJECTED", b.pk))
        Site.objects.create(site_id="XXX_1", name="NEW")
        rl = P.relaunch(b, "Site XXX_1 créé dans Gestion des sites")
        self.assertEqual((rl.relaunch_of_id, rl.rows_duplicate, FuelStockObservation.objects.count()), (b.pk, 1, 2))


class WorkQueueTests(TmpRootMixin, TestCase):
    def setUp(self):
        super().setUp()
        for sid in ("BKL_0086", "DKR_0002", "DKR_0003"):
            Site.objects.create(site_id=sid, name=sid)

    def test_queue_messages_are_actionable(self):
        P.process_bytes(xlsx({"K": [HEAD, guard_row("BKL_0086", cm=24)]}), "a.xlsx", "LOLIFORCE", origin="UPLOAD")
        tank("DKR_0002", status="A_CONFIRMER", family=None, dims=(None, None, None), plate=5000)
        tank("DKR_0003", status="INCOMPLETE", family=None, dims=(None, None, None), plate=None)
        c = {x["key"]: x["count"] for x in W.counters()}
        self.assertEqual((c["releves_non_convertibles"], c["cuves_a_confirmer"], c["cuves_incompletes"], c["gardiennage_a_configurer"]), (1, 1, 1, 1))
        item = W.build("releves_non_convertibles")[0]
        self.assertEqual(item["message"], "BKL_0086 — Relevé Loliforce : 24 cm — aucune cuve validée pour ce site. "
                                          "Action : renseigner ou confirmer la cuve avant conversion.")
        self.assertEqual((item["action_type"], item["regle"]), ("OPEN_TANK", "Conversion cm → L : cuve site VALIDÉE requise"))
        tank_item = W.build("cuves_a_confirmer")[0]
        self.assertIn("plaque 5000 L", tank_item["valeur_brute"])
        self.assertTrue(tank_item["action"])


class AuditAndApiTests(TmpRootMixin, TestCase):
    def setUp(self):
        super().setUp()
        Site.objects.create(site_id="DKR_0001", name="ALPHA", installed_typology="A_Ax_GG S2", installed_site_type="OUTDOOR")
        User = get_user_model()
        self.manager = User.objects.create_user("mgr", "mgr@x.sn", "p", role="manager")
        self.analyst = User.objects.create_user("ana", "ana@x.sn", "p", role="analyst")
        self.client = APIClient()

    def test_tank_correction_validation_reconverts_and_is_traced(self):
        t = tank("DKR_0001", status="PROPOSEE", dims=(100, None, 100), plate=1000)
        P.process_bytes(xlsx({"K": [HEAD, guard_row("DKR_0001", cm=10)]}), "a.xlsx", "LOLIFORCE", origin="UPLOAD")
        o = FuelStockObservation.objects.get()
        self.assertEqual(o.issue_codes[0], "TANK_PROPOSED")
        self.client.force_authenticate(self.analyst)
        self.assertEqual(self.client.patch(f"/api/fuel-tracking/stock/cuves/{t.pk}/", {"dimension_2_cm": 100, "reason": "x"}, format="json").status_code, 403)
        self.client.force_authenticate(self.manager)
        r = self.client.patch(f"/api/fuel-tracking/stock/cuves/{t.pk}/", {"validate": True, "reason": "Contrôle"}, format="json")
        self.assertEqual(r.status_code, 400)  # hauteur manquante
        r = self.client.patch(f"/api/fuel-tracking/stock/cuves/{t.pk}/", {"dimension_2_cm": 100, "validate": True, "reason": "Mesure terrain 14/10"}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual((r.json()["validation_status"], r.json()["plate_coherent"]), ("VALIDEE", True))
        o.refresh_from_db()
        self.assertEqual((float(o.level_liters), o.conversion_method, o.validation_status), (100.0, "CALIBRATED", "PENDING"))  # gardiennage non configuré
        c = FuelRegistryCorrection.objects.get(entity="TANK", field="dimension_2_cm")
        self.assertEqual((c.old_value, c.new_value, c.reason, c.author), (None, "100", "Mesure terrain 14/10", self.manager))
        self.assertEqual(self.client.post(f"/api/fuel-tracking/stock/observations/{o.pk}/valider/", {}, format="json").status_code, 400)
        r = self.client.post(f"/api/fuel-tracking/stock/observations/{o.pk}/valider/", {"reason": "Gardien confirmé"}, format="json")
        self.assertEqual(r.json()["statut_releve"], "VALIDE_MANUELLEMENT")
        row = self.client.get("/api/fuel-tracking/stock/registre/?search=dkr_0001").json()["data"][0]
        self.assertEqual((row["stock_l"], row["source"], row["conversion_method"], row["statut"], row["typologie"], row["environnement"]),
                         (100.0, "LOLIFORCE", "CALIBRATED", "RETENU", "A_Ax_GG S2", "Outdoor"))
        self.assertEqual(row["profil_energetique"], "Incomplet / à confirmer")  # jamais déduit de la typologie
        hist = self.client.get("/api/fuel-tracking/stock/registre/DKR_0001/").json()
        self.assertEqual(hist["observations"][0]["volume_calibrated_l"], 100.0)
        self.assertTrue(hist["corrections"])
        q = self.client.get("/api/fuel-tracking/stock/a-traiter/").json()
        self.assertTrue(any(x["key"] == "cuves_a_confirmer" for x in q["counters"]))

    def test_last_validated_vs_last_known_movement(self):
        FuelMovement.objects.create(source_type="HISTORICAL_REFUELING", source_event_id="x", site_id="DKR_0001", movement_type="ADD",
                                    movement_datetime=timezone.now(), status="CLOSED")
        self.client.force_authenticate(self.analyst)
        row = self.client.get("/api/fuel-tracking/stock/registre/?search=DKR_0001").json()["data"][0]
        self.assertIsNone(row["last_validated_movement"])
        self.assertEqual(row["last_movement"]["status"], "CLOSED")

    def test_next_action_and_action_filter(self):
        P.process_bytes(xlsx({"K": [HEAD, guard_row("DKR_0001", cm=10)]}), "a.xlsx", "LOLIFORCE", origin="UPLOAD")
        self.client.force_authenticate(self.analyst)
        r = self.client.get("/api/fuel-tracking/stock/registre/?action_requise=1").json()
        row = next(x for x in r["data"] if x["site_id"] == "DKR_0001")
        self.assertEqual(row["statut"], "INDISPONIBLE")
        self.assertTrue(row["prochaine_action"])  # sans cuve validée : action explicite, jamais un simple « à valider »
        self.assertEqual(r["kpis"]["sites_action_requise"], r["pagination"]["total"])
        none = self.client.get("/api/fuel-tracking/stock/registre/?statut=AUCUNE_DONNEE").json()["data"]
        self.assertTrue(all(x["prochaine_action"] is None for x in none))
        ranked = self.client.get("/api/fuel-tracking/stock/registre/?tri=priorite&limit=200").json()["data"]
        self.assertEqual(ranked[0]["site_id"], "DKR_0001")  # relevé présent avant les sites sans relevé

    def test_upload_models_and_enoc_types_api(self):
        self.client.force_authenticate(self.manager)
        r = self.client.post("/api/fuel-tracking/stock/imports/", {"file": SimpleUploadedFile("steve.xlsx", steve_workbook()), "source_type": "TANK_MODELS"})
        self.assertEqual(r.status_code, 201, r.content)
        models = self.client.get("/api/fuel-tracking/stock/modeles-cuves/").json()
        self.assertEqual(len(models["data"]), 8)
        r = self.client.post("/api/fuel-tracking/stock/enoc/types/", {"raw_value": "AJOUT", "category": "AJOUT", "reason": "Ops"}, format="json")
        self.assertEqual((r.status_code, r.json()["version"]), (201, 1))
        self.client.force_authenticate(self.analyst)
        self.assertEqual(self.client.post("/api/fuel-tracking/stock/enoc/types/", {"raw_value": "X", "category": "AJOUT", "reason": "x"}, format="json").status_code, 403)
        self.assertEqual(self.client.get("/api/fuel-tracking/stock/audit-profil-energetique/").status_code, 200)

    def test_audit_command_runs_every_section_read_only(self):
        from django.core.management import call_command
        tank("DKR_0001")
        P.process_bytes(xlsx({"K": [HEAD, guard_row("DKR_0001", cm=10)]}), "a.xlsx", "LOLIFORCE", origin="UPLOAD")
        before = (FuelStockObservation.objects.count(), FuelRegistryCorrection.objects.count(), FuelTankMaster.objects.count())
        out = io.StringIO()
        call_command("fuel_stock_audit", stdout=out)
        text = out.getvalue()
        self.assertNotIn("Section en erreur", text)
        for title in ("1. Bibliothèque", "4. Stock retenu", "9. Cas de bout en bout"):
            self.assertIn(title, text)
        self.assertEqual(before, (FuelStockObservation.objects.count(), FuelRegistryCorrection.objects.count(),
                                  FuelTankMaster.objects.count()))  # lecture seule

    def test_guard_scope_command_reevaluates(self):
        from django.core.management import call_command
        tank("DKR_0001")
        P.process_bytes(xlsx({"K": [HEAD, guard_row("DKR_0001", cm=10)]}), "a.xlsx", "LOLIFORCE", origin="UPLOAD")
        path = Path(tempfile.mkdtemp()) / "perimetre.csv"
        path.write_text("site_id;provider;active_from\ndkr_0001;LOLIFORCE;01/01/2026\n", encoding="utf-8")
        out = io.StringIO()
        call_command("import_fuel_guard_scope", file=str(path), stdout=out)
        self.assertIn("1 périmètre(s) créé(s)", out.getvalue())
        o = FuelStockObservation.objects.get()
        self.assertEqual((o.guard_scope_status, o.validation_status), ("IN_SCOPE", "VALIDATED"))

    def test_guard_scope_clears_not_configured_even_if_reading_stays_pending(self):
        from django.core.management import call_command
        tank("DKR_0001", status="PROPOSEE")
        P.process_bytes(xlsx({"K": [HEAD, guard_row("DKR_0001", cm=10)]}), "a.xlsx", "LOLIFORCE", origin="UPLOAD")
        self.assertIn("NOT_CONFIGURED", FuelStockObservation.objects.get().issue_codes)
        path = Path(tempfile.mkdtemp()) / "p.csv"
        path.write_text("site_id;provider;active_from\nDKR_0001;LOLIFORCE;01/01/2026\n", encoding="utf-8")
        for _ in range(2):  # relancer l'import réévalue sans dupliquer le périmètre
            call_command("import_fuel_guard_scope", file=str(path), stdout=io.StringIO())
        o = FuelStockObservation.objects.get()
        self.assertEqual((o.guard_scope_status, o.validation_status), ("IN_SCOPE", "PENDING"))  # cuve seulement proposée
        self.assertNotIn("NOT_CONFIGURED", o.issue_codes)
        self.assertEqual(FuelGuardScope.objects.count(), 1)

    def test_guard_scope_api_reevaluates(self):
        tank("DKR_0001")
        P.process_bytes(xlsx({"K": [HEAD, guard_row("DKR_0001", cm=10)]}), "a.xlsx", "LOLIFORCE", origin="UPLOAD")
        self.client.force_authenticate(self.manager)
        r = self.client.post("/api/fuel-tracking/stock/gardiennage/",
                             {"site_id": "dkr_0001", "provider": "LOLIFORCE", "active_from": "01/01/2026", "reason": "Contrat 2026"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        o = FuelStockObservation.objects.get()
        self.assertEqual((o.guard_scope_status, o.validation_status, R.display_status(o)), ("IN_SCOPE", "VALIDATED", "AUTO_VALIDE"))
