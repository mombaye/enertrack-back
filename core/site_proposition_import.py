"""
Import du fichier Ops « Proposition_de_load_MMYYYY.xlsx » (feuille Export, en-têtes en ligne 2)
dans le référentiel commun des sites (core.Site), depuis Gestion des sites → Import Sites.

Règles :
  - un site absent est créé ; un site existant est retrouvé sans tenir compte de la casse
    (DkR_0123 = DKR_0123) ;
  - seules les valeurs renseignées du fichier sont écrites : une cellule vide n'efface jamais
    une valeur existante ;
  - colonnes reprises : Site_Name → name, Typologie facturée → billing_typology,
    Typologies cible Contractuel → ordered_typology, Typologie réelle → installed_typology,
    Configuration (Indoor/Outdoor) → site_type / installed_site_type ;
  - Average_Load du mois (colonne Month) → load mensuel du module financier (SiteMonthlyLoad,
    source Prévisionnel), comme l'import des loads d'Évaluation financière.
Les autres modules (Suivi Carburant, financier, facturation…) lisent core.Site : les sites
créés ou modifiés y apparaissent sans autre action.
"""
from __future__ import annotations

import io

import openpyxl
from django.db import transaction

from .models import Site

ZONES = {code for code, _ in Site.ZONE_CHOICES}
FIELDS = {
    "site_name": "name",
    "typologie facturée": "billing_typology",
    "typologies cible contractuel": "ordered_typology",
    "typologie réelle": "installed_typology",
}


def is_proposition_file(file_bytes: bytes) -> bool:
    """Fichier Proposition de load : une ligne d'en-têtes avec Site_ID et Average_Load (ou Typologie réelle)."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(min_row=1, max_row=5, values_only=True))
        wb.close()
    except Exception:
        return False
    for row in rows:
        heads = {str(v).strip().lower() for v in row if v is not None}
        if "site_id" in heads and ({"average_load", "typologie réelle"} & heads):
            return True
    return False


def _clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _kind(v):
    s = (_clean(v) or "").lower().replace(" ", "").replace("-", "")
    return {"indoor": "INDOOR", "outdoor": "OUTDOOR"}.get(s)


def _normalize_id(site_id: str) -> str:
    prefix, sep, rest = site_id.partition("_")
    return f"{prefix.upper()}{sep}{rest}" if sep else site_id.upper()


def _read_rows(file_bytes: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    header, out = None, []
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if any(isinstance(v, str) and v.strip().lower() == "site_id" for v in row):
                header = [str(v).strip().lower() if v is not None else None for v in row]
            continue
        # Première occurrence d'un en-tête (le fichier répète « Date du relevé(PM+CM) »).
        rec = {}
        for i, h in enumerate(header):
            if h and h not in rec and i < len(row):
                rec[h] = row[i]
        if _clean(rec.get("site_id")):
            out.append(rec)
    wb.close()
    return out


def import_proposition(file_bytes: bytes, user) -> dict:
    from financial.models import SiteMonthlyLoad
    from financial.views import _parse_load_proposition

    rows = _read_rows(file_bytes)
    by_upper = {s.site_id.upper(): s for s in Site.objects.all()}
    country = getattr(user, "pays", None) or "sen"
    created = updated = unchanged = 0
    normalized_ids: list[str] = []
    errors: list[dict] = []
    id_map: dict[str, str] = {}

    with transaction.atomic():
        for n, rec in enumerate(rows, start=3):
            raw_id = _clean(rec.get("site_id"))
            try:
                site = by_upper.get(raw_id.upper())
                values = {field: _clean(rec.get(col)) for col, field in FIELDS.items()}
                kind = _kind(rec.get("configuration"))
                if site is None:
                    site_id = _normalize_id(raw_id)
                    if site_id != raw_id:
                        normalized_ids.append(f"{raw_id} → {site_id}")
                    prefix = site_id.split("_")[0]
                    site = Site(site_id=site_id, country=country, zone=prefix if prefix in ZONES else None)
                    is_new = True
                else:
                    is_new = False
                id_map[raw_id] = site.site_id
                changed = []
                for field, value in values.items():
                    if value is not None and getattr(site, field) != value:
                        setattr(site, field, value)
                        changed.append(field)
                if kind:
                    for field in ("site_type", "installed_site_type"):
                        if getattr(site, field) != kind:
                            setattr(site, field, kind)
                            changed.append(field)
                if not site.configuration and site.installed_typology:
                    site.configuration = site.installed_typology
                    changed.append("configuration")
                if changed or is_new:
                    parts = [p for p in ((site.configuration or "").strip(), (site.site_type or "").strip(),
                                         (site.load_band or "").strip()) if p]
                    site.target_mapping_key = " | ".join(parts) if parts else None
                    site.save()
                    by_upper[site.site_id.upper()] = site
                    if is_new:
                        created += 1
                    else:
                        updated += 1
                else:
                    unchanged += 1
            except Exception as e:  # une ligne en erreur n'arrête pas l'import
                errors.append({"row": n, "site_id": raw_id, "error": str(e)})

        # Load mensuel (Average_Load) → module financier, sur les identifiants du référentiel.
        loads = {"created": 0, "updated": 0, "month": None}
        sites = {s.site_id: s for s in Site.objects.filter(site_id__in=set(id_map.values()))}
        for r in _parse_load_proposition(file_bytes):
            site = sites.get(id_map.get(r["site_id"], r["site_id"]))
            if site is None:
                continue
            _, was_created = SiteMonthlyLoad.objects.update_or_create(
                site=site, year=r["year"], month=r["month"],
                defaults={"load_w": r["load_w"], "source": SiteMonthlyLoad.Source.PREVISIONNEL, "imported_by": user},
            )
            loads["created" if was_created else "updated"] += 1
            loads["month"] = f"{r['year']}-{r['month']:02d}"

    return {
        "message": (f"Import Proposition de load terminé : {created} site(s) créé(s), {updated} mis à jour, "
                    f"{unchanged} inchangé(s) ; load du mois {loads['month'] or '—'} : "
                    f"{loads['created'] + loads['updated']} site(s)."),
        "format": "proposition_de_load",
        "rows": len(rows),
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
        "normalized_ids": normalized_ids,
        "loads": loads,
        "errors_count": len(errors),
        "errors": errors[:50],
    }
